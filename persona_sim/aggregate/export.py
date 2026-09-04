"""出力ファイル一式の書き出し（`SPEC_PHASE1.md` §7.4）。

```
outputs/{survey_id}/
  ├ crosstab_{measure}.csv   … measure ごと。コンセプトを表側、選択肢を表頭に置く
  ├ crosstab_all.csv         … 全設問を1枚に積んだもの
  ├ panel_composition.csv
  ├ open_ends.csv
  ├ responses_raw.csv
  ├ run_metadata.json        … `persona_sim.metadata` が run 時に書く
  └ report.xlsx              … 上の表を1ブックにまとめたもの
```

CSV は BOM 付き UTF-8 で書く。Excel が既定のエンコーディングで開くと日本語が壊れるため。
"""

from __future__ import annotations

import csv
import re
from collections.abc import Iterable, Sequence
from pathlib import Path
from typing import TYPE_CHECKING, Any

from persona_sim import attribution_text
from persona_sim.aggregate import frame
from persona_sim.aggregate.tables import Table
from persona_sim.config import StorageConfig
from persona_sim.errors import PersonaSimError
from persona_sim.panel.schema import SurveyDefinition

if TYPE_CHECKING:  # pragma: no cover
    from pyspark.sql import SparkSession

    from persona_sim.aggregate.aggregate import AggregateResult

#: Excel で開いたときに日本語が化けないようにする。
CSV_ENCODING = "utf-8-sig"

#: xlsx のシート名に使えない文字（Excel の制約）。
_INVALID_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")
_MAX_SHEET_NAME = 31

FORMAT_CSV = "csv"
FORMAT_XLSX = "xlsx"


def write_outputs(
    spark: SparkSession,
    survey: SurveyDefinition,
    storage: StorageConfig,
    result: AggregateResult,
    output_dir: str,
    *,
    formats: Sequence[str] | None = None,
) -> list[Path]:
    """§7.4 の一式を書き出す。書いたファイルのパスを返す。"""
    selected = tuple(formats if formats is not None else survey.output.formats)
    destination = Path(output_dir) / survey.survey_id
    destination.mkdir(parents=True, exist_ok=True)

    written: list[Path] = []
    if FORMAT_CSV in selected:
        for table in result.tables:
            written.append(write_csv(destination / f"{table.key}.csv", table.columns, table.rows))
        written.append(_write_open_ends(spark, survey, storage, destination))
        written.append(_write_responses_raw(spark, survey, storage, destination))

    if FORMAT_XLSX in selected:
        # ローデータは responses_raw.csv にあるので、ブックには積まない
        # （Excel の行数上限に当たるうえ、同じものを2つ配ることになる）。
        written.append(
            write_workbook(
                destination / "report.xlsx",
                survey,
                result.tables,
                notes=result.notes,
            )
        )

    return written


def write_csv(path: Path, columns: Sequence[str], rows: Iterable[Sequence[Any]]) -> Path:
    with open(path, "w", encoding=CSV_ENCODING, newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(columns)
        writer.writerows(rows)
    return path


def _write_open_ends(
    spark: SparkSession, survey: SurveyDefinition, storage: StorageConfig, destination: Path
) -> Path:
    columns, rows = frame.load_open_ends(spark, survey, storage)
    return write_csv(destination / "open_ends.csv", columns, rows)


def _write_responses_raw(
    spark: SparkSession, survey: SurveyDefinition, storage: StorageConfig, destination: Path
) -> Path:
    columns, rows = frame.stream_responses_raw(spark, survey, storage)
    return write_csv(destination / "responses_raw.csv", columns, rows)


# --------------------------------------------------------------------------- #
# xlsx
#
# **writer は1本だけ。** バッチ実行の `report.xlsx`（§7.4）と Web UI のダウンロード
# （`docs/SPEC_UI.md` §4.4）は、渡す表とローデータの有無が違うだけで構成は同じにする。
# 別実装にしていた頃は、片方に入れた注記や帰属表示がもう片方から落ちていた。
# --------------------------------------------------------------------------- #

#: シート名。
SUMMARY_SHEET = "概要"
RAW_SHEET = "ローデータ"

#: Excel の1シートあたりの行数上限。
EXCEL_MAX_ROWS = 1_048_576


def write_workbook(
    path: Path,
    survey: SurveyDefinition,
    tables: Sequence[Table],
    *,
    notes: Sequence[str] = (),
    raw_columns: Sequence[str] = (),
    raw_rows: Sequence[Sequence[Any]] = (),
) -> Path:
    """表を1ブックにまとめる。

    先頭に「概要」シートを置き、調査情報・注記（§11 の E3・E4）・読み方・帰属表示を
    そこに集める。**表だけ配ると注記が読まれずに終わる**ため、シートを分けても
    必ず1枚目に載る位置に置く。

    `raw_columns` / `raw_rows` を渡すと「ローデータ」シートを足す（UI のダウンロードは
    1ファイルしか返せないので、集計表と同じブックに入れる）。
    """
    workbook = _new_workbook()
    summary = workbook.active
    summary.title = SUMMARY_SHEET
    for row in _summary_rows(survey, tables, notes):
        summary.append(row)

    used: set[str] = {summary.title}
    for table in tables:
        sheet = workbook.create_sheet(_sheet_name(table, used))
        sheet.append([table.title])
        sheet.append([])
        sheet.append(list(table.columns))
        for row in table.rows:
            sheet.append(list(row))
        if table.notes:
            sheet.append([])
            for note in table.notes:
                sheet.append([f"注記: {note}"])

    if raw_columns:
        raw = workbook.create_sheet(RAW_SHEET)
        raw.append(list(raw_columns))
        truncated = _truncate(raw_rows)
        for row in truncated:
            raw.append(list(row))
        if len(truncated) < len(raw_rows):
            # 黙って切らない。切られたことに気づかないまま「全件」として配られる方が害が大きい。
            raw.append(
                [
                    f"※ Excel の行数上限のため {len(truncated):,} 件で打ち切った"
                    f"（全 {len(raw_rows):,} 件）。全件が要る場合は"
                    " responses_raw.csv を使うこと"
                ]
            )

    workbook.save(path)
    return path


def _new_workbook():
    try:
        from openpyxl import Workbook
    except ImportError as exc:  # pragma: no cover - 依存が入っていない環境向け
        raise PersonaSimError(
            "report.xlsx の書き出しには openpyxl が必要。"
            "`pip install openpyxl` するか、survey の output.formats から xlsx を外すこと"
        ) from exc
    return Workbook()


def _summary_rows(
    survey: SurveyDefinition, tables: Sequence[Table], notes: Sequence[str]
) -> list[list[Any]]:
    rows: list[list[Any]] = [
        ["調査ID", survey.survey_id],
        ["調査名", survey.name],
        ["コンセプト数", len(survey.stimuli)],
        ["設問数", len(survey.questions)],
        ["セグメント軸", ", ".join(survey.output.segments)],
        [],
        ["シート"],
    ]
    rows.extend([table.title] for table in tables)
    rows.extend(
        [
            [],
            ["注記"],
        ]
    )
    rows.extend([[note] for note in (notes or ["なし"])])
    rows.extend(
        [
            [],
            ["読み方"],
            ["n はウェイト適用前の実数、% はウェイト適用後（ウェイトが全て1.0なら一致する）"],
            ["品質フラグの立った回答も集計に含めている。除いた場合の n を併記してある"],
            ["平均は選択肢番号を逆順スコア化した値（順序尺度の設問のみ）"],
            ["コンセプト間の相対比較として読むこと。統計的推測の指標は算出していない"],
        ]
    )
    # 帰属表示（§15.3）と免責（§15.1）。**表だけ配られる前提で必ず載せる。**
    rows.append([])
    rows.extend([line] for line in attribution_text().splitlines())
    return rows


def _truncate(rows: Sequence[Sequence[Any]]) -> Sequence[Sequence[Any]]:
    """見出し1行と打ち切り注記1行のぶんを残して収める。"""
    limit = EXCEL_MAX_ROWS - 2
    return rows if len(rows) <= limit else rows[:limit]


def _sheet_name(table: Table, used: set[str]) -> str:
    """Excel のシート名制約（31文字・記号不可・重複不可）に収める。"""
    base = _INVALID_SHEET_CHARS.sub("_", table.key)[:_MAX_SHEET_NAME] or "sheet"
    name = base
    suffix = 2
    while name in used:
        tail = f"_{suffix}"
        name = base[: _MAX_SHEET_NAME - len(tail)] + tail
        suffix += 1
    used.add(name)
    return name
